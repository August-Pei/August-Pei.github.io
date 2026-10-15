#!/usr/bin/env python3
"""快手原始 CSV 导入、账号文件命名、总量写入和格式校验。"""
import argparse
import csv
from datetime import date
from decimal import Decimal, InvalidOperation
import json
import os
from pathlib import Path
import shutil
import sys
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data/kuaishou'
BASE = ['作品标题', '发布时间', '播放量', '评论量', '点赞量']
TOTALS = ['收藏量', '转发量', '粉丝量']


def text(value):
    return str(value).strip() if value is not None else ''


def count(value, label):
    if isinstance(value, bool) or value is None or text(value) == '':
        raise ValueError(f'{label} 必须为非负整数')
    raw = text(value).replace(',', '').replace('，', '')
    multiplier = 10000 if raw[-1:].lower() in ('万', 'w') else 1
    if multiplier != 1:
        raw = raw[:-1]
    try:
        number = Decimal(raw) * multiplier
        if not number.is_finite() or number != number.to_integral_value() or not 0 <= number <= 9007199254740991:
            raise ValueError()
        return int(number)
    except (InvalidOperation, ValueError):
        raise ValueError(f'{label} 不是有效的非负整数：{value}')


def read_csv(file):
    with file.open(encoding='utf-8-sig', newline='') as handle:
        rows = list(csv.reader(handle, strict=True))
    while rows and not any(rows[-1]):
        rows.pop()
    return rows


def video_info(rows, complete=False, expected=None):
    if len(rows) < 2:
        raise ValueError('视频 CSV 没有数据')
    header = [text(value) for value in rows[0]]
    for key in BASE:
        if header.count(key) != 1:
            raise ValueError(f'视频表必须有唯一的「{key}」列')
    for key in TOTALS:
        if header.count(key) > 1:
            raise ValueError(f'存在重复的「{key}」列')
    for i, row in enumerate(rows[1:], 1):
        if len(row) != len(header):
            raise ValueError(f'CSV 第 {i + 1} 行列数与表头不一致')
    totals = [i for i, row in enumerate(rows[1:], 1) if row and text(row[header.index('作品标题')]) == '总计']
    if len(totals) > 1 or totals and totals[0] != len(rows) - 1:
        raise ValueError('总计行必须唯一且位于最后一行')
    details = [(i, row) for i, row in enumerate(rows[1:], 1) if i not in totals and any(row)]
    if not details:
        raise ValueError('缺少视频明细')
    for _, row in details:
        for key in BASE[2:]:
            count(row[header.index(key)], key)
    result = {'rows': len(details), 'header': header, 'totalRow': totals[0] + 1 if totals else None}
    if complete:
        if header[-3:] != TOTALS or not totals:
            raise ValueError('末尾必须为收藏量、转发量、粉丝量三列，且最后一行必须为总计，请先写入三项总量')
        values = {}
        for key in TOTALS:
            raw = rows[totals[0]][header.index(key)]
            if not text(raw).isascii() or not text(raw).isdigit():
                raise ValueError(f'总计「{key}」必须为完整的非负整数')
            values[key] = count(raw, key)
            if expected is not None and values[key] != expected[key]:
                raise ValueError(f'总计「{key}」为 {values[key]}，与输入的 {expected[key]} 不一致')
        result['totals'] = values
    return result


def validate(file, kind, complete=False, expected=None):
    if kind == 'videos':
        return video_info(read_csv(file), complete, expected)
    from openpyxl import load_workbook
    book = load_workbook(file, read_only=True, data_only=True)
    try:
        rows = list(book.active.iter_rows(values_only=True))
        if len(rows) < 2:
            raise ValueError('账号数据表没有数据')
        header = [text(value) for value in rows[0]]
        for aliases in [('日期', '时间', 'date'), ('播放量', 'views', 'view'), ('净增粉丝量', '粉丝净增', '净增粉丝', '净涨粉')]:
            if not any(key in header for key in aliases):
                raise ValueError(f'账号数据缺少「{aliases[0]}」列，请导出帐号数据分析明细表')
        return {'rows': sum(any(value is not None for value in row) for row in rows[1:])}
    finally:
        book.close()


def backup(file):
    if not file.exists():
        return None
    folder = ROOT / 'data/admin-backups/kuaishou' / file.parent.name
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f'{file.stem}-{uuid4().hex}{file.suffix}'
    shutil.copy2(file, target)
    return str(target)


def import_file(folder, source, kind):
    validate(source, kind)
    target = folder / ('videos.csv' if kind == 'videos' else 'account.xlsx')
    temporary = folder / f'.{kind}-{uuid4().hex}{target.suffix}'
    try:
        shutil.copy2(source, temporary)
        backup(target)
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return [f'已导入并命名为 {target.name}，文件内容保持原样。']


def normalize(folder):
    matches = [file for file in folder.iterdir() if file.is_file() and file.suffix.lower() == '.xlsx' and file.name.startswith(('帐号数据分析明细表', '账号数据分析明细表'))]
    if len(matches) > 1 or matches and (folder / 'account.xlsx').exists():
        raise ValueError('帐号文件存在多个候选或 account.xlsx 已存在，请先保留一份待重命名文件')
    if not matches:
        return ['没有待重命名的帐号数据分析明细表；可直接上传下载文件并命名为 account.xlsx。']
    validate(matches[0], 'account')
    name = matches[0].name
    matches[0].rename(folder / 'account.xlsx')
    return [f'{name} → account.xlsx']


def update(folder, expected):
    file = folder / 'videos.csv'
    rows = read_csv(file)
    info = video_info(rows)
    header = info['header']
    existing = [key for key in header if key in TOTALS]
    if existing and (existing != TOTALS[:len(existing)] or header[-len(existing):] != existing):
        raise ValueError('已有总量列必须位于末尾并按收藏量、转发量、粉丝量排列')
    extra = TOTALS[len(existing):]
    rows[0].extend(extra)
    for row in rows[1:]:
        row.extend([''] * len(extra))
    header = [text(value) for value in rows[0]]
    if info['totalRow']:
        total_row = rows[info['totalRow'] - 1]
    else:
        total_row = [''] * len(header)
        total_row[header.index('作品标题')] = '总计'
        rows.append(total_row)
    for key in TOTALS:
        total_row[header.index(key)] = str(expected[key])
    temporary = folder / f'.videos-{uuid4().hex}.csv'
    try:
        with temporary.open('w', encoding='utf-8-sig', newline='') as handle:
            csv.writer(handle).writerows(rows)
        validate(temporary, 'videos', complete=True, expected=expected)
        saved_backup = backup(file)
        os.replace(temporary, file)
        record_temp = folder / f'.acquisition-{uuid4().hex}.tmp'
        try:
            record_temp.write_text(json.dumps(expected, ensure_ascii=False), encoding='utf-8')
            os.replace(record_temp, folder / '.acquisition.json')
        finally:
            record_temp.unlink(missing_ok=True)
        return [f'已备份原 CSV：{saved_backup}', f'已写入第 {len(rows)} 行总计：' + '，'.join(f'{key} {expected[key]}' for key in TOTALS), '已重新读取 CSV，三列、总计行与三项数值检查通过。原有视频字段保持原样。']
    finally:
        temporary.unlink(missing_ok=True)


def check(expected=None):
    folders = []
    if DATA.exists():
        for folder in DATA.iterdir():
            try:
                if folder.is_dir() and date.fromisoformat(folder.name).isoformat() == folder.name:
                    folders.append(folder)
            except ValueError:
                pass
    folder = max(folders, key=lambda folder: folder.name) if folders else None
    if folder and expected is None and (folder / '.acquisition.json').exists():
        expected = json.loads((folder / '.acquisition.json').read_text(encoding='utf-8'))
    files = []
    if folder:
        for kind, name in [('videos', 'videos.csv'), ('account', 'account.xlsx')]:
            file = folder / name
            result = {'name': name, 'exists': file.is_file(), 'valid': False}
            try:
                if not result['exists']:
                    raise ValueError('文件不存在')
                result.update(validate(file, kind, complete=kind == 'videos', expected=expected))
                if kind == 'videos' and (folder / 'videos.xlsx').exists():
                    raise ValueError('目录中另有 videos.xlsx，汇总脚本会优先读取它；请移出该旧文件后再使用 CSV 数据')
                result.update(valid=True, bytes=file.stat().st_size, message='三项总量末列、总计末行和数值检查通过' if kind == 'videos' else '账号数据可以读取，必需表头存在')
            except Exception as error:
                result['message'] = str(error)
            files.append(result)
    return {'date': folder.name if folder else None, 'directory': str(folder or DATA), 'expected': expected, 'files': files, 'valid': len(files) == 2 and all(file['valid'] for file in files)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action', choices=('check', 'import', 'normalize', 'update'))
    parser.add_argument('--date')
    parser.add_argument('--file', type=Path)
    parser.add_argument('--kind', choices=('videos', 'account'))
    parser.add_argument('--totals')
    args = parser.parse_args()
    try:
        expected = None
        if args.totals is not None:
            values = json.loads(args.totals)
            expected = {key: count(values[key], key) for key in TOTALS}
        if args.action == 'check':
            result = check(expected)
        else:
            if not args.date or date.fromisoformat(args.date).isoformat() != args.date:
                raise ValueError('日期格式不正确')
            folder = DATA / args.date
            if not folder.is_dir():
                raise ValueError('请先创建当天目录')
            if args.action == 'update':
                if expected is None:
                    raise ValueError('请填写三项总量')
                logs = update(folder, expected)
            elif args.action == 'normalize':
                logs = normalize(folder)
            else:
                logs = import_file(folder, args.file, args.kind)
            result = {'logs': logs, 'check': check(expected)}
        print(json.dumps(result, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({'error': str(error)}, ensure_ascii=False))
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())

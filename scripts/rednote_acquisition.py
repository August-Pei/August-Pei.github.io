#!/usr/bin/env python3
"""后台小红书导出文件的校验、命名与粉丝数写入（不访问网络）。"""
import argparse
from copy import copy
from datetime import date
import json
import math
import os
from pathlib import Path
import shutil
import sys
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data' / 'rednote'
KINDS = ('notes', 'views', 'engagement', 'fans')
METRICS = ('浏览量', '点赞量', '收藏量', '评论数', '转发数')


def text(value):
    return str(value).strip() if value is not None else ''


def integer(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or int(value) != value:
        raise ValueError('粉丝数必须为非负整数')
    if not 0 <= value <= 9007199254740991:
        raise ValueError('粉丝数超出有效范围')
    return int(value)


def rows_of(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    while rows and all(value is None for value in rows[-1]):
        rows.pop()
    return rows


def notes_structure(sheet, complete=False, expected=None):
    rows = rows_of(sheet)
    if len(rows) < 2:
        raise ValueError('笔记表为空或没有笔记数据')
    header = [text(value) for value in rows[0]]
    while header and not header[-1]:
        header.pop()
    for key in ('笔记ID', *METRICS):
        if header.count(key) != 1:
            raise ValueError(f'笔记表必须有唯一的「{key}」列')
    fans_columns = [i for i, value in enumerate(header) if value == '粉丝数']
    if len(fans_columns) > 1:
        raise ValueError('存在重复的「粉丝数」列，请先清理')
    totals = [i for i, row in enumerate(rows[1:], 1) if any(text(value) == '总计' for value in row)]
    if len(totals) > 1 or totals and totals[0] != len(rows) - 1:
        raise ValueError('「总计」行必须唯一且位于最后一行')
    id_column = header.index('笔记ID')
    details = [row for i, row in enumerate(rows[1:], 1) if i not in totals and text(row[id_column])]
    if not details:
        raise ValueError('笔记表缺少有效的笔记明细')
    fans = None
    if complete:
        if len(fans_columns) != 1 or fans_columns[0] != len(header) - 1:
            raise ValueError('最后一列必须为「粉丝数」')
        if not totals:
            raise ValueError('缺少末尾「总计」行，请先写入粉丝数')
        fans = integer(rows[totals[0]][fans_columns[0]])
        if expected is not None and fans != expected:
            raise ValueError(f'总计粉丝数为 {fans}，与输入的 {expected} 不一致')
    return {'rows': len(details), 'fans': fans, 'header': header, 'totalRow': totals[0] + 1 if totals else None, 'lastRow': len(rows)}


def validate_book(file, kind, complete=False, expected=None):
    from openpyxl import load_workbook
    book = load_workbook(file, read_only=True, data_only=True)
    try:
        if kind == 'notes':
            return notes_structure(book.active, complete, expected)
        if kind in ('views', 'fans'):
            sheet_name = '观看趋势' if kind == 'views' else '净涨粉趋势'
            if sheet_name not in book.sheetnames:
                raise ValueError(f'缺少「{sheet_name}」工作表，请下载完整的 30 日数据')
            rows = rows_of(book[sheet_name])
            header = [text(value) for value in rows[0]] if rows else []
            date_columns = [i for i, value in enumerate(header) if value in ('日期', '时间', 'date')]
            aliases = ('数值', '观看量', '播放量', 'views') if kind == 'views' else ('数值', '净涨粉', '净增粉丝', '粉丝净增')
            value_columns = [i for i, value in enumerate(header) if value in aliases]
            if not date_columns or not value_columns:
                raise ValueError(f'「{sheet_name}」工作表缺少日期或数值列')
            values = [row[value_columns[0]] for row in rows[1:] if text(row[date_columns[0]])]
            if not values or any(value is None or isinstance(value, bool) or not math.isfinite(float(value)) for value in values):
                raise ValueError(f'「{sheet_name}」工作表没有有效趋势数据')
            return {'rows': len(values), 'sheets': book.sheetnames}
        rows = rows_of(book.active)
        if len(rows) < 2 or not any(any(value is not None for value in row) for row in rows[1:]):
            raise ValueError('互动数据表为空或没有数据')
        header = [text(value) for value in rows[0]]
        if '指标' not in header or '数值' not in header:
            raise ValueError('互动数据表缺少「指标」「数值」列，请下载完整互动数据')
        return {'rows': len(rows) - 1, 'sheets': book.sheetnames}
    finally:
        book.close()


def check(expected=None):
    directories = []
    if DATA.exists():
        for folder in DATA.iterdir():
            try:
                if folder.is_dir() and date.fromisoformat(folder.name).isoformat() == folder.name:
                    directories.append(folder)
            except ValueError:
                pass
    folder = max(directories, key=lambda item: item.name) if directories else None
    if folder and expected is None:
        record = folder / '.acquisition.json'
        if record.exists():
            expected = integer(json.loads(record.read_text())['fans'])
    files = []
    if folder:
        for kind in KINDS:
            file = folder / f'{kind}.xlsx'
            result = {'name': file.name, 'exists': file.is_file(), 'valid': False}
            try:
                if not result['exists']:
                    raise ValueError('文件不存在')
                result.update(validate_book(file, kind, complete=kind == 'notes', expected=expected))
                result.update(valid=True, bytes=file.stat().st_size, message='格式正确' if kind != 'notes' else '粉丝数末列、总计末行及粉丝数校验通过')
            except Exception as error:
                result['message'] = str(error)
            files.append(result)
    return {'date': folder.name if folder else None, 'directory': str(folder or DATA), 'expectedFans': expected, 'files': files, 'valid': len(files) == 4 and all(file['valid'] for file in files)}


def backup(file):
    destination = ROOT / 'data' / 'admin-backups' / 'rednote' / file.parent.name
    destination.mkdir(parents=True, exist_ok=True)
    target = destination / f'{file.stem}-{uuid4().hex}.xlsx'
    shutil.copy2(file, target)
    return str(target)


def update_notes(folder, fans):
    from openpyxl import load_workbook
    file = folder / 'notes.xlsx'
    validate_book(file, 'notes')
    book = load_workbook(file)
    temporary = folder / f'.notes-{uuid4().hex}.xlsx'
    try:
        sheet = book.active
        info = notes_structure(sheet)
        header = info['header']
        last_column = len(header)
        while last_column and not header[last_column - 1]:
            last_column -= 1
        if '粉丝数' in header:
            fans_column = header.index('粉丝数') + 1
            if fans_column != last_column:
                raise ValueError('已有「粉丝数」列不在最后一列，请先调整原表')
        else:
            fans_column = last_column + 1
            cell = sheet.cell(1, fans_column, '粉丝数')
            cell._style = copy(sheet.cell(1, max(1, last_column))._style)
            from openpyxl.utils import get_column_letter
            sheet.column_dimensions[get_column_letter(fans_column)].width = 14
        total_row = info['totalRow'] or info['lastRow'] + 1
        if not info['totalRow']:
            for column in range(1, fans_column + 1):
                sheet.cell(total_row, column)._style = copy(sheet.cell(max(1, total_row - 1), min(column, max(1, last_column)))._style)
            sheet.cell(total_row, header.index('笔记ID') + 1, '总计')
        sheet.cell(total_row, fans_column, fans).number_format = '0'
        book.save(temporary)
        validate_book(temporary, 'notes', complete=True, expected=fans)
        saved_backup = backup(file)
        os.replace(temporary, file)
        record = folder / '.acquisition.json'
        record_temp = folder / f'.acquisition-{uuid4().hex}.tmp'
        try:
            record_temp.write_text(json.dumps({'fans': fans}, ensure_ascii=False), encoding='utf-8')
            os.replace(record_temp, record)
        finally:
            record_temp.unlink(missing_ok=True)
        return [f'已备份原表：{saved_backup}', f'已写入 notes.xlsx：第 {fans_column} 列「粉丝数」，第 {total_row} 行「总计」，粉丝数 {fans}', '已重新打开保存文件，格式与粉丝数校验通过']
    finally:
        book.close()
        temporary.unlink(missing_ok=True)


def normalize(folder):
    prefixes = {'notes': '小红书创作者全部笔记', 'views': '近30日观看数据', 'engagement': '近30日互动数据', 'fans': '近30日涨粉数据'}
    plans = []
    for kind, prefix in prefixes.items():
        target = folder / f'{kind}.xlsx'
        matches = [file for file in folder.iterdir() if file.is_file() and file.suffix.lower() == '.xlsx' and file.name.startswith(prefix)]
        if len(matches) > 1 or matches and target.exists():
            raise ValueError(f'{kind}.xlsx 存在重复候选文件，请保留一份待导入文件后重试')
        if matches:
            validate_book(matches[0], kind)
            plans.append((matches[0], target))
    logs = []
    for source, target in plans:
        source.rename(target)
        logs.append(f'{source.name} → {target.name}')
    return logs or ['没有待重命名的下载文件；可上传文件，或将原始下载文件放入当前目录再重试。']


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action', choices=('check', 'validate', 'normalize', 'update'))
    parser.add_argument('--date')
    parser.add_argument('--file', type=Path)
    parser.add_argument('--kind', choices=KINDS)
    parser.add_argument('--fans', type=int)
    args = parser.parse_args()
    try:
        if args.fans is not None:
            integer(args.fans)
        if args.action == 'check':
            result = check(args.fans)
        elif args.action == 'validate':
            result = validate_book(args.file, args.kind)
        else:
            if not args.date or date.fromisoformat(args.date).isoformat() != args.date:
                raise ValueError('日期格式不正确')
            folder = DATA / args.date
            if not folder.is_dir():
                raise ValueError('请先创建当天目录')
            if args.action == 'update' and args.fans is None:
                raise ValueError('请填写粉丝数')
            logs = update_notes(folder, args.fans) if args.action == 'update' else normalize(folder)
            result = {'logs': logs, 'check': check(args.fans)}
        print(json.dumps(result, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({'error': str(error)}, ensure_ascii=False))
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())

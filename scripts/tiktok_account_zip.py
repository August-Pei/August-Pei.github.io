#!/usr/bin/env python3
"""将 TikTok Studio ZIP 内的 Overview.xlsx 导入为 account.xlsx。"""
import argparse
import hashlib
import json
import os
from pathlib import Path, PurePosixPath
import shutil
import stat
import sys
from uuid import uuid4
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
DOWNLOADS = Path('/Users/pyh/Downloads')


def digest(file):
    value = hashlib.sha256()
    with file.open('rb') as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b''):
            value.update(chunk)
    return value.digest()


def import_zip(archive, download_name=None):
    from openpyxl import load_workbook
    if archive.is_symlink() or not archive.is_file():
        raise ValueError('ZIP 文件不存在或不是普通文件')
    folder = archive.parent
    if folder.parent.resolve() != (ROOT / 'data/TikTok').resolve():
        raise ValueError('ZIP 必须位于 TikTok 日期目录中')
    target = folder / 'account.xlsx'
    temporary = folder / f'.account-{uuid4().hex}.xlsx'
    original_hash = digest(archive) if download_name else None
    logs = []
    try:
        with ZipFile(archive) as zipped:
            entries = zipped.infolist()
            if len(entries) > 1000 or sum(item.file_size for item in entries) > 128 * 1024 * 1024:
                raise ValueError('ZIP 展开数据过大')
            candidates = []
            for item in entries:
                route = PurePosixPath(item.filename)
                if route.is_absolute() or '..' in route.parts or '\\' in item.filename or stat.S_ISLNK(item.external_attr >> 16) or item.flag_bits & 1:
                    raise ValueError('ZIP 包含不支持的路径、链接或加密文件')
                if item.is_dir() or '__MACOSX' in route.parts or route.name.startswith('._'):
                    continue
                if route.name.lower() == 'overview.xlsx':
                    candidates.append(item)
            if len(candidates) != 1:
                raise ValueError('ZIP 中必须有且只有一份 Overview.xlsx')
            with zipped.open(candidates[0]) as incoming, temporary.open('xb') as outgoing:
                shutil.copyfileobj(incoming, outgoing)
        book = load_workbook(temporary, read_only=True, data_only=True)
        try:
            if not book.sheetnames or not any(value is not None for row in book.active.iter_rows(min_row=1, max_row=2, values_only=True) for value in row):
                raise ValueError('Overview.xlsx 为空')
        finally:
            book.close()
        if target.exists():
            backup = ROOT / 'data/admin-backups/TikTok' / folder.name
            backup.mkdir(parents=True, exist_ok=True)
            shutil.copy2(target, backup / f'account-{uuid4().hex}.xlsx')
        os.replace(temporary, target)
        logs.append('已解压 Overview.xlsx，确认 Excel 可读取，重命名为 account.xlsx。')
        try:
            archive.unlink()
            logs.append('已删除本次导入的 ZIP 文件。')
        except OSError as error:
            logs.append(f'account.xlsx 已导入，但 ZIP 清理失败，请手动删除 {archive.name}：{error}')
        if download_name and Path(download_name).name == download_name and download_name.startswith('Overview') and download_name.lower().endswith('.zip'):
            downloaded = DOWNLOADS / download_name
            try:
                if downloaded.is_file() and not downloaded.is_symlink():
                    if digest(downloaded) == original_hash:
                        downloaded.unlink()
                        logs.append(f'已删除 Downloads 中同名且内容一致的原 ZIP：{download_name}')
                    else:
                        logs.append('Downloads 同名 ZIP 内容不同，已保留该文件。')
            except OSError as error:
                logs.append(f'account.xlsx 已导入，Downloads 原 ZIP 未能清理：{error}')
        return {'logs': logs}
    finally:
        temporary.unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('archive', type=Path)
    parser.add_argument('--download-name')
    args = parser.parse_args()
    try:
        print(json.dumps(import_zip(args.archive, args.download_name), ensure_ascii=False))
    except Exception as error:
        print(json.dumps({'error': str(error)}, ensure_ascii=False))
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())

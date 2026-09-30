#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""growth.py：计算各平台最近 30 个统计日的每日存量并输出 JSON。

平台后台导出的趋势文件保存的是每日增量。本脚本使用
``overall_data.py`` 读取到的最新粉丝和播放量存量作为锚点，再按日期
倒推历史存量。

日期约定：

* 国内平台：下载文件夹日期的前一天是最后统计日；
* YouTube / TikTok：原始数据日期顺延一天后与国内统计日对齐；
* Instagram 没有相关趋势数据，不参与统计。
"""

from __future__ import annotations

import csv
import json
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Callable, Iterable

from openpyxl import load_workbook

from overall_data import (
    BILIBILI_DIR,
    DOUYIN_DIR,
    KUAISHOU_DIR,
    REDNOTE_DIR,
    TIKTOK_DIR,
    YOUTUBE_DIR,
    find_latest_dir,
    read_bilibili_data,
    read_douyin_data,
    read_kuaishou_data,
    read_rednote_data,
    read_tiktok_data,
    read_youtube_data,
    to_int,
)


DAYS = 30
OVERSEAS_PLATFORMS = {"youtube", "tiktok"}


def normalize_header(value: Any) -> str:
    """统一表头格式，便于兼容空格、大小写及 BOM。"""
    return re.sub(r"\s+", " ", str(value or "").lstrip("\ufeff").strip()).casefold()


def find_header_index(header: Iterable[Any], aliases: Iterable[str]) -> int | None:
    indexes = {normalize_header(value): index for index, value in enumerate(header)}
    for alias in aliases:
        index = indexes.get(normalize_header(alias))
        if index is not None:
            return index
    return None


def parse_date(value: Any, reference: date) -> date | None:
    """解析平台导出的常见日期格式。

    对不含年份的 ``M月D日``，选择最靠近当前统计末日的年份。这样也能
    正确处理跨年 30 日窗口。
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value or "").strip()
    if not text:
        return None

    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y年%m月%d日"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass

    match = re.fullmatch(r"(\d{1,2})月(\d{1,2})日", text)
    if match:
        month, day = map(int, match.groups())
        candidates = []
        for year in (reference.year - 1, reference.year, reference.year + 1):
            try:
                candidate = date(year, month, day)
            except ValueError:
                continue
            candidates.append(candidate)
        if candidates:
            return min(candidates, key=lambda candidate: abs(candidate - reference))

    return None


def read_csv(path: Path) -> list[tuple[Any, ...]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        return [tuple(row) for row in csv.reader(file)]


def read_xlsx(path: Path, sheet_name: str | None = None) -> list[tuple[Any, ...]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise KeyError(f"{path} 缺少工作表 {sheet_name!r}")
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.active
        return [tuple(row) for row in worksheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def read_increment_table(
    path: Path,
    *,
    reference: date,
    date_aliases: tuple[str, ...],
    value_aliases: tuple[str, ...],
    sheet_name: str | None = None,
    shift_days: int = 0,
) -> dict[date, int]:
    """读取两列式或多列式表格，返回对齐后的 ``日期 -> 当日增量``。"""
    rows = read_csv(path) if path.suffix.lower() == ".csv" else read_xlsx(path, sheet_name)
    if not rows:
        raise ValueError(f"数据文件为空: {path}")

    date_index = find_header_index(rows[0], date_aliases)
    value_index = find_header_index(rows[0], value_aliases)
    if date_index is None:
        raise KeyError(f"{path} 缺少日期列（支持: {', '.join(date_aliases)}）")
    if value_index is None:
        raise KeyError(f"{path} 缺少数值列（支持: {', '.join(value_aliases)}）")

    increments: dict[date, int] = {}
    for row in rows[1:]:
        if date_index >= len(row) or value_index >= len(row):
            continue
        source_day = parse_date(row[date_index], reference - timedelta(days=shift_days))
        if source_day is None:
            continue
        aligned_day = source_day + timedelta(days=shift_days)
        increments[aligned_day] = increments.get(aligned_day, 0) + to_int(row[value_index])
    return increments


def latest_stats_day(root: Path) -> tuple[Path, date]:
    latest_dir = find_latest_dir(root)
    download_day = datetime.strptime(latest_dir.name, "%Y-%m-%d").date()
    return latest_dir, download_day - timedelta(days=1)


def find_file(directory: Path, filenames: tuple[str, ...]) -> Path:
    for filename in filenames:
        path = directory / filename
        if path.is_file():
            return path
    raise FileNotFoundError(f"{directory} 中未找到: {', '.join(filenames)}")


def read_bilibili_increments(directory: Path, end_day: date) -> tuple[dict[date, int], Path]:
    path = find_file(directory, ("account.csv",))
    return read_increment_table(
        path,
        reference=end_day,
        date_aliases=("时间", "日期", "date"),
        value_aliases=("净增粉丝", "粉丝净增", "net followers"),
    ), path


def read_rednote_increments(directory: Path, end_day: date) -> tuple[dict[date, int], Path]:
    path = find_file(directory, ("fans.xlsx", "fans.csv"))
    return read_increment_table(
        path,
        reference=end_day,
        date_aliases=("日期", "时间", "date"),
        value_aliases=("数值", "净涨粉", "净增粉丝", "粉丝净增"),
        sheet_name="净涨粉趋势" if path.suffix.lower() == ".xlsx" else None,
    ), path


def read_douyin_increments(directory: Path, end_day: date) -> tuple[dict[date, int], Path]:
    path = find_file(directory, ("fans.xlsx", "fans.csv"))
    return read_increment_table(
        path,
        reference=end_day,
        date_aliases=("日期", "时间", "date"),
        value_aliases=("粉丝净增", "净增粉丝", "净涨粉"),
    ), path


def read_kuaishou_increments(directory: Path, end_day: date) -> tuple[dict[date, int], Path]:
    path = find_file(directory, ("account.xlsx", "account.csv", "fans.xlsx", "fans.csv"))
    return read_increment_table(
        path,
        reference=end_day,
        date_aliases=("日期", "时间", "date"),
        value_aliases=("净增粉丝量", "粉丝净增", "净增粉丝", "净涨粉"),
    ), path


def read_youtube_increments(directory: Path, end_day: date) -> tuple[dict[date, int], Path]:
    path = find_file(directory, ("account.csv", "account.xlsx", "fans.csv", "fans.xlsx"))
    return read_increment_table(
        path,
        reference=end_day,
        date_aliases=("Date", "日期", "时间"),
        value_aliases=("Subscribers gained", "Net subscribers", "粉丝净增", "净增粉丝"),
        shift_days=1,
    ), path


def read_tiktok_increments(directory: Path, end_day: date) -> tuple[dict[date, int], Path]:
    path = find_file(directory, ("account.xlsx", "account.csv", "fans.xlsx", "fans.csv"))
    return read_increment_table(
        path,
        reference=end_day,
        date_aliases=("Date", "日期", "时间"),
        value_aliases=(
            "Followers gained",
            "Net followers",
            "New followers",
            "Followers",
            "粉丝净增",
            "净增粉丝",
            "新增粉丝",
            "净涨粉",
        ),
        shift_days=1,
    ), path


def read_bilibili_view_increments(directory: Path, end_day: date) -> tuple[dict[date, int], Path]:
    path = find_file(directory, ("account.csv",))
    return read_increment_table(
        path,
        reference=end_day,
        date_aliases=("时间", "日期", "date"),
        value_aliases=("播放量", "views", "view"),
    ), path


def read_rednote_view_increments(directory: Path, end_day: date) -> tuple[dict[date, int], Path]:
    path = find_file(directory, ("views.xlsx", "views.csv"))
    return read_increment_table(
        path,
        reference=end_day,
        date_aliases=("日期", "时间", "date"),
        value_aliases=("数值", "观看量", "播放量", "views"),
        sheet_name="观看趋势" if path.suffix.lower() == ".xlsx" else None,
    ), path


def read_douyin_view_increments(directory: Path, end_day: date) -> tuple[dict[date, int], Path]:
    path = find_file(directory, ("views & engagement.xlsx", "views & engagement.csv"))
    return read_increment_table(
        path,
        reference=end_day,
        date_aliases=("日期", "时间", "date"),
        value_aliases=("总播放量", "播放量", "views"),
    ), path


def read_kuaishou_view_increments(directory: Path, end_day: date) -> tuple[dict[date, int], Path]:
    path = find_file(directory, ("account.xlsx", "account.csv"))
    return read_increment_table(
        path,
        reference=end_day,
        date_aliases=("日期", "时间", "date"),
        value_aliases=("播放量", "views", "view"),
    ), path


def read_youtube_view_increments(directory: Path, end_day: date) -> tuple[dict[date, int], Path]:
    path = find_file(directory, ("account.csv", "account.xlsx"))
    return read_increment_table(
        path,
        reference=end_day,
        date_aliases=("Date", "日期", "时间"),
        value_aliases=("Views", "View", "播放量", "观看次数"),
        shift_days=1,
    ), path


def read_tiktok_view_increments(directory: Path, end_day: date) -> tuple[dict[date, int], Path]:
    path = find_file(directory, ("account.xlsx", "account.csv"))
    return read_increment_table(
        path,
        reference=end_day,
        date_aliases=("Date", "日期", "时间"),
        value_aliases=("Video Views", "Views", "播放量", "观看次数"),
        shift_days=1,
    ), path


IncrementReader = Callable[[Path, date], tuple[dict[date, int], Path]]
CurrentReader = Callable[[], dict[str, Any]]

PLATFORMS: tuple[
    tuple[str, Path, IncrementReader, IncrementReader, CurrentReader], ...
] = (
    (
        "bilibili",
        BILIBILI_DIR,
        read_bilibili_increments,
        read_bilibili_view_increments,
        read_bilibili_data,
    ),
    (
        "rednote",
        REDNOTE_DIR,
        read_rednote_increments,
        read_rednote_view_increments,
        read_rednote_data,
    ),
    (
        "douyin",
        DOUYIN_DIR,
        read_douyin_increments,
        read_douyin_view_increments,
        read_douyin_data,
    ),
    (
        "kuaishou",
        KUAISHOU_DIR,
        read_kuaishou_increments,
        read_kuaishou_view_increments,
        read_kuaishou_data,
    ),
    (
        "youtube",
        YOUTUBE_DIR,
        read_youtube_increments,
        read_youtube_view_increments,
        read_youtube_data,
    ),
    (
        "tiktok",
        TIKTOK_DIR,
        read_tiktok_increments,
        read_tiktok_view_increments,
        read_tiktok_data,
    ),
)


def build_history(current_fans: int, end_day: date, increments: dict[date, int]) -> list[dict[str, Any]]:
    days = [end_day - timedelta(days=offset) for offset in reversed(range(DAYS))]
    stock_by_day: dict[date, int] = {}
    running_stock = current_fans

    for day in reversed(days):
        stock_by_day[day] = running_stock
        running_stock -= increments.get(day, 0)

    return [{"date": day.isoformat(), "fans": stock_by_day[day]} for day in days]


def build_view_history(current_views: int, end_day: date, increments: dict[date, int]) -> list[dict[str, Any]]:
    """以最后一日的累计播放量为锚点，逐日扣除当日增量向前倒推。"""
    days = [end_day - timedelta(days=offset) for offset in reversed(range(DAYS))]
    stock_by_day: dict[date, int] = {}
    running_stock = current_views

    for day in reversed(days):
        stock_by_day[day] = running_stock
        running_stock -= increments.get(day, 0)

    return [{"date": day.isoformat(), "views": stock_by_day[day]} for day in days]


def build_data() -> dict[str, Any]:
    platform_results: dict[str, Any] = {}
    warnings: list[str] = []

    for key, root, fan_increment_reader, view_increment_reader, current_reader in PLATFORMS:
        latest_dir, end_day = latest_stats_day(root)
        current_data = current_reader()
        current_fans = to_int(current_data.get("fans"))
        current_views = to_int(current_data.get("views"))
        fallback_used = False

        try:
            fan_increments, source_path = fan_increment_reader(latest_dir, end_day)
        except KeyError as error:
            # 有最新存量、但导出表不含粉丝趋势时，仍输出可消费的 30 日结构；
            # 元数据和 warnings 会明确标记该平台没有可用于倒推的增量。
            if key != "tiktok":
                raise
            source_path = find_file(
                latest_dir,
                ("account.xlsx", "account.csv", "fans.xlsx", "fans.csv"),
            )
            fan_increments = {}
            fallback_used = True
            warnings.append(f"tiktok: {error}；30 日内暂按 0 增量计算")

        view_increments, view_source_path = view_increment_reader(latest_dir, end_day)
        fan_history = build_history(current_fans, end_day, fan_increments)
        view_history = build_view_history(current_views, end_day, view_increments)
        history = [
            {
                "date": fan_point["date"],
                "fans": fan_point["fans"],
                "views": view_point["views"],
            }
            for fan_point, view_point in zip(fan_history, view_history, strict=True)
        ]
        platform_results[key] = {
            "download_date": latest_dir.name,
            "start_date": history[0]["date"],
            "end_date": history[-1]["date"],
            "source": str(source_path.relative_to(source_path.parents[2])),
            "views_source": str(view_source_path.relative_to(view_source_path.parents[2])),
            "overseas_date_shift_days": 1 if key in OVERSEAS_PLATFORMS else 0,
            "increment_data_available": not fallback_used,
            "history": history,
        }

    # 正常情况下各平台来自同一天下载的数据。仅在 30 日区间完全一致时才
    # 计算全网存量，避免把不同统计日的数字强行相加。
    ranges = {
        (item["start_date"], item["end_date"])
        for item in platform_results.values()
    }
    total_history: list[dict[str, Any]] | None = None
    if len(ranges) == 1:
        total_by_date: dict[str, dict[str, int]] = {}
        for item in platform_results.values():
            for point in item["history"]:
                totals = total_by_date.setdefault(point["date"], {"fans": 0, "views": 0})
                totals["fans"] += point["fans"]
                totals["views"] += point["views"]
        total_history = [
            {"date": day, **totals}
            for day, totals in sorted(total_by_date.items())
        ]
    else:
        warnings.append("各平台最新统计日不一致，未生成全网存量合计")

    return {
        "days": DAYS,
        "platforms": platform_results,
        "total": total_history,
        "warnings": warnings,
    }


def main() -> None:
    try:
        print(json.dumps(build_data(), ensure_ascii=False, indent=2))
    except (FileNotFoundError, KeyError, ValueError) as error:
        print(f"增长数据读取失败: {error}", file=sys.stderr)
        raise SystemExit(1) from error


if __name__ == "__main__":
    main()

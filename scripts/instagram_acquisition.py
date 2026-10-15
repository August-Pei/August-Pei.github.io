#!/usr/bin/env python3
"""SuperIns 文件导入及 Instagram shares / fans 总量写入。"""
import argparse
from copy import copy
import csv
from datetime import date
import json
import math
import os
from pathlib import Path
import shutil
import sys
from uuid import uuid4

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / 'data/Instagram'
KEYS = ('shares', 'fans')


def integer(value, key):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or int(value) != value or not 0 <= value <= 9007199254740991:
        raise ValueError(f'{key} 必须为有效的非负整数')
    return int(value)


def rows_of(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    while rows and all(value is None for value in rows[-1]):
        rows.pop()
    return rows


def structure(sheet, complete=False, expected=None):
    rows = rows_of(sheet)
    if len(rows) < 2:
        raise ValueError('videos.xlsx 没有视频数据')
    header = [str(value).strip() if value is not None else '' for value in rows[0]]
    while header and not header[-1]:
        header.pop()
    for key in ('views', 'likes', 'comments'):
        if header.count(key) != 1:
            raise ValueError(f'必须有唯一的 {key} 列，请检查 SuperIns 导出设置')
    for key in KEYS:
        if header.count(key) > 1:
            raise ValueError(f'存在重复的 {key} 列')
    totals = [i for i, row in enumerate(rows[1:], 1) if row and str(row[0]).strip() == 'All']
    if len(totals) > 1 or totals and totals[0] != len(rows) - 1:
        raise ValueError('All 行必须唯一且位于最后一行')
    details = [row for i, row in enumerate(rows[1:], 1) if i not in totals and any(value is not None for value in row)]
    if not details:
        raise ValueError('缺少视频明细')
    result = {'rows': len(details), 'header': header, 'lastRow': len(rows), 'totalRow': totals[0] + 1 if totals else None}
    if complete:
        if header[-2:] != list(KEYS) or not totals:
            raise ValueError('请先写入 shares、fans 两个末列及最后的 All 行')
        values = {key: integer(rows[totals[0]][header.index(key)], key) for key in KEYS}
        if expected is not None:
            for key in KEYS:
                if values[key] != expected[key]:
                    raise ValueError(f'All 行 {key} 为 {values[key]}，与输入的 {expected[key]} 不一致')
        result['totals'] = values
    return result


def validate(file, complete=False, expected=None):
    from openpyxl import load_workbook
    book = load_workbook(file, read_only=True, data_only=True)
    try:
        return structure(book.active, complete, expected)
    finally:
        book.close()


def backup(file):
    if not file.exists():
        return None
    folder = ROOT / 'data/admin-backups/Instagram' / file.parent.name
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f'{file.stem}-{uuid4().hex}{file.suffix}'
    shutil.copy2(file, target)
    return str(target)


def import_file(folder, source):
    from openpyxl import Workbook
    temporary = folder / f'.videos-{uuid4().hex}.xlsx'
    book = None
    try:
        if source.suffix.lower() == '.csv':
            with source.open(encoding='utf-8-sig', newline='') as handle:
                rows = list(csv.reader(handle, strict=True))
            if not rows:
                raise ValueError('SuperIns CSV 为空')
            header = [value.strip() for value in rows[0]]
            book = Workbook()
            sheet = book.active
            sheet.title = 'Instagram'
            for r, row in enumerate(rows, 1):
                if len(row) != len(header):
                    raise ValueError(f'CSV 第 {r} 行列数与表头不一致')
                for c, value in enumerate(row, 1):
                    cell = sheet.cell(r, c)
                    if r > 1 and header[c - 1] in ('views', 'likes', 'comments', *KEYS) and value.strip():
                        number = float(value.strip().replace(',', ''))
                        cell.value = integer(number, header[c - 1])
                    else:
                        cell.value = value or None
                        if value:
                            cell.data_type = 's'
            sheet.freeze_panes = 'A2'
            book.save(temporary)
        else:
            shutil.copy2(source, temporary)
        validate(temporary)
        target = folder / 'videos.xlsx'
        backup(target)
        os.replace(temporary, target)
        return ['已导入 videos.xlsx' + ('（由 CSV 转换，视频字段和文本内容保留）' if source.suffix.lower() == '.csv' else '，保留原有工作表与数据')]
    finally:
        if book:
            book.close()
        temporary.unlink(missing_ok=True)


def normalize(folder):
    candidates = [file for file in folder.iterdir() if file.is_file() and file.name.startswith('SuperIns_') and file.suffix.lower() in ('.xlsx', '.csv')]
    if len(candidates) > 1 or candidates and (folder / 'videos.xlsx').exists():
        raise ValueError('存在多个 SuperIns 导出文件或 videos.xlsx 已存在，请先保留一份待导入文件')
    if not candidates:
        return ['没有待命名的 SuperIns 文件；请上传原始导出文件，或放入当天目录后重试。']
    source = candidates[0]
    if source.suffix.lower() == '.csv':
        return import_file(folder, source)
    validate(source)
    source.rename(folder / 'videos.xlsx')
    return [f'{source.name} → videos.xlsx']


def update(folder, expected):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    file = folder / 'videos.xlsx'
    validate(file)
    book = load_workbook(file)
    temporary = folder / f'.videos-{uuid4().hex}.xlsx'
    try:
        sheet = book.active
        info = structure(sheet)
        header = info['header']
        existing = [key for key in header if key in KEYS]
        if existing and (existing != list(KEYS[:len(existing)]) or header[-len(existing):] != existing):
            raise ValueError('已有 shares / fans 列须位于末尾，并按 shares、fans 顺序排列')
        for key in KEYS[len(existing):]:
            previous = len(header)
            header.append(key)
            cell = sheet.cell(1, len(header), key)
            cell._style = copy(sheet.cell(1, previous)._style)
            sheet.column_dimensions[get_column_letter(len(header))].width = 16
        total_row = info['totalRow'] or info['lastRow'] + 1
        if not info['totalRow']:
            sheet.cell(total_row, 1, 'All')
        for key in KEYS:
            sheet.cell(total_row, header.index(key) + 1, expected[key]).number_format = '0'
        book.save(temporary)
        validate(temporary, complete=True, expected=expected)
        saved_backup = backup(file)
        os.replace(temporary, file)
        record_temp = folder / f'.acquisition-{uuid4().hex}.tmp'
        try:
            record_temp.write_text(json.dumps(expected), encoding='utf-8')
            os.replace(record_temp, folder / '.acquisition.json')
        finally:
            record_temp.unlink(missing_ok=True)
        return [f'原表已备份：{saved_backup}', f'已写入第 {total_row} 行 All：shares={expected["shares"]}，fans={expected["fans"]}', '已重新读取保存文件，末尾两列、All 行及输入数值检查通过。']
    finally:
        book.close()
        temporary.unlink(missing_ok=True)


def check(expected=None):
    folders = []
    if DATA.exists():
        for folder in DATA.iterdir():
            try:
                if folder.is_dir() and date.fromisoformat(folder.name).isoformat() == folder.name:
                    folders.append(folder)
            except ValueError:
                pass
    if not folders:
        return {'date': None, 'directory': str(DATA), 'files': [], 'expected': expected, 'valid': False}
    folder = max(folders, key=lambda item: item.name)
    if expected is None and (folder / '.acquisition.json').exists():
        expected = json.loads((folder / '.acquisition.json').read_text(encoding='utf-8'))
    file = folder / 'videos.xlsx'
    item = {'name': file.name, 'exists': file.is_file(), 'valid': False}
    try:
        if not item['exists']:
            raise ValueError('文件不存在')
        item.update(validate(file, complete=True, expected=expected))
        item.update(valid=True, bytes=file.stat().st_size, message='shares / fans 末列、唯一 All 末行及两项总量检查通过')
    except Exception as error:
        item['message'] = str(error)
    return {'date': folder.name, 'directory': str(folder), 'expected': expected, 'files': [item], 'valid': item['valid']}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('action', choices=('check', 'import', 'normalize', 'update'))
    parser.add_argument('--date')
    parser.add_argument('--file', type=Path)
    parser.add_argument('--totals')
    args = parser.parse_args()
    try:
        expected = None
        if args.totals is not None:
            values = json.loads(args.totals)
            expected = {key: integer(values[key], key) for key in KEYS}
        if args.action == 'check':
            result = check(expected)
        else:
            if not args.date or date.fromisoformat(args.date).isoformat() != args.date:
                raise ValueError('日期格式不正确')
            folder = DATA / args.date
            if not folder.is_dir():
                raise ValueError('请先创建当天目录')
            if args.action == 'update':
                if expected is None:
                    raise ValueError('请填写 shares 和 fans')
                logs = update(folder, expected)
            elif args.action == 'normalize':
                logs = normalize(folder)
            else:
                logs = import_file(folder, args.file)
            result = {'logs': logs, 'check': check(expected)}
        print(json.dumps(result, ensure_ascii=False))
    except Exception as error:
        print(json.dumps({'error': str(error)}, ensure_ascii=False))
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
